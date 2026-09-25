import os
from datetime import datetime, timedelta
from io import BytesIO

from bson import ObjectId
from dotenv import load_dotenv
from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pymongo import ASCENDING, DESCENDING, MongoClient
from pymongo.errors import PyMongoError

# ---------------------------------------------------------
# App / environment setup
# ---------------------------------------------------------
BASE_DIR = os.path.abspath(os.path.dirname(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change-this-secret-key")

ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin123")
MONGO_URI = os.getenv("MONGO_URI", "").strip()
MONGO_DB_NAME = os.getenv("MONGO_DB_NAME", "VisitorDB")
MONGO_COLLECTION_NAME = os.getenv("MONGO_COLLECTION_NAME", "visitors")

if not MONGO_URI:
    raise RuntimeError(
        "MONGO_URI is not configured. "
        "Create a .env file from .env.example and paste your MongoDB Atlas connection string."
    )

mongo_client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=5000)
db = mongo_client[MONGO_DB_NAME]
visitors_collection = db[MONGO_COLLECTION_NAME]


def verify_mongodb():
    """Fail fast if MongoDB cannot be reached."""
    mongo_client.admin.command("ping")


def ensure_indexes():
    """Small indexes that help the visitor/admin screens."""
    visitors_collection.create_index([("check_in", DESCENDING)])
    visitors_collection.create_index([("check_out", ASCENDING)])
    visitors_collection.create_index([("visitor_name", ASCENDING)])
    visitors_collection.create_index([("company", ASCENDING)])


# ---------------------------------------------------------
# Helper functions
# ---------------------------------------------------------
def admin_required():
    return session.get("admin") is True


def to_object_id(value):
    try:
        return ObjectId(value)
    except Exception:
        return None


def display_datetime(value):
    """Admin UI: date + time without seconds."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return "-"


def export_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return ""


def export_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    return ""


def serialize_visitor(doc):
    """Prepare MongoDB document for Jinja templates."""
    return {
        "id": str(doc["_id"]),
        "company": doc.get("company", ""),
        "visitor_name": doc.get("visitor_name", ""),
        "host_name": doc.get("host_name", ""),
        "host_phone": doc.get("host_phone", ""),
        "purpose": doc.get("purpose", ""),
        "comment": doc.get("comment", ""),
        "checkout_comment": doc.get("checkout_comment", ""),
        "safety_agreed": bool(doc.get("safety_agreed", False)),
        "check_in": doc.get("check_in"),
        "check_out": doc.get("check_out"),
        "check_in_display": display_datetime(doc.get("check_in")),
        "check_out_display": display_datetime(doc.get("check_out")),
    }


def build_admin_filter(q="", date_from="", date_to="", status=""):
    filters = []

    if q:
        filters.append(
            {
                "$or": [
                    {"company": {"$regex": q, "$options": "i"}},
                    {"visitor_name": {"$regex": q, "$options": "i"}},
                    {"host_name": {"$regex": q, "$options": "i"}},
                    {"host_phone": {"$regex": q, "$options": "i"}},
                    {"purpose": {"$regex": q, "$options": "i"}},
                    {"comment": {"$regex": q, "$options": "i"}},
                    {"checkout_comment": {"$regex": q, "$options": "i"}},
                ]
            }
        )

    if date_from:
        try:
            start = datetime.strptime(date_from, "%Y-%m-%d")
            filters.append({"check_in": {"$gte": start}})
        except ValueError:
            pass

    if date_to:
        try:
            end = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
            filters.append({"check_in": {"$lt": end}})
        except ValueError:
            pass

    if status == "inside":
        filters.append({"check_out": None})
    elif status == "checked_out":
        filters.append({"check_out": {"$ne": None}})

    if not filters:
        return {}

    if len(filters) == 1:
        return filters[0]

    return {"$and": filters}


@app.context_processor
def inject_helpers():
    return {
        "current_year": datetime.now().year,
    }


# ---------------------------------------------------------
# Visitor check-in
# ---------------------------------------------------------
@app.route("/", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        company = request.form.get("company", "").strip()
        visitor_name = request.form.get("visitor_name", "").strip()
        host_name = request.form.get("host_name", "").strip()
        host_phone = request.form.get("host_phone", "").strip()
        purpose = request.form.get("purpose", "").strip()
        comment = request.form.get("comment", "").strip()
        safety_agreed = request.form.get("safety_agreed") == "on"

        if not company or not visitor_name or not host_name or not purpose:
            flash("Please complete all required fields.", "error")
            return render_template("register.html")

        if not safety_agreed:
            flash(
                "You must agree to the Safety Guidelines before check-in.",
                "error",
            )
            return render_template("register.html")

        now = datetime.now()

        try:
            visitors_collection.insert_one(
                {
                    "company": company,
                    "visitor_name": visitor_name,
                    "host_name": host_name,
                    "host_phone": host_phone,
                    "purpose": purpose,
                    "comment": comment,
                    "safety_agreed": True,
                    "check_in": now,
                    "check_out": None,
                }
            )
        except PyMongoError as exc:
            app.logger.exception("MongoDB insert failed")
            flash(f"Could not save the visitor record: {exc}", "error")
            return render_template("register.html")

        return render_template(
            "success.html",
            visitor_name=visitor_name,
            check_in=now.strftime("%Y-%m-%d %H:%M"),
        )

    return render_template("register.html")


# ---------------------------------------------------------
# Visitor self-service check-out
# ---------------------------------------------------------
@app.route("/checkout", methods=["GET", "POST"])
def visitor_checkout():
    query = request.form.get("query", "").strip() if request.method == "POST" else ""

    mongo_filter = {"check_out": None}

    if query:
        mongo_filter["$or"] = [
            {"visitor_name": {"$regex": query, "$options": "i"}},
            {"company": {"$regex": query, "$options": "i"}},
        ]

    try:
        docs = list(
            visitors_collection.find(mongo_filter)
            .sort("check_in", DESCENDING)
            .limit(100)
        )
    except PyMongoError as exc:
        app.logger.exception("MongoDB checkout search failed")
        flash(f"Could not load visitor records: {exc}", "error")
        docs = []

    results = [serialize_visitor(doc) for doc in docs]

    if request.method == "POST" and query and not results:
        flash("No active visitor record was found.", "error")

    return render_template(
        "visitor_checkout.html",
        results=results,
        query=query,
    )


@app.post("/checkout/confirm/<visitor_id>")
def visitor_checkout_confirm(visitor_id):
    object_id = to_object_id(visitor_id)

    if not object_id:
        flash("Invalid visitor record.", "error")
        return redirect(url_for("visitor_checkout"))

    try:
        visitor = visitors_collection.find_one({"_id": object_id, "check_out": None})

        if not visitor:
            flash(
                "This visitor has already checked out or the record was not found.",
                "error",
            )
            return redirect(url_for("visitor_checkout"))

        now = datetime.now()
        checkout_comment = request.form.get("checkout_comment", "").strip()[:1000]

        result = visitors_collection.update_one(
            {"_id": object_id, "check_out": None},
            {"$set": {
                "check_out": now,
                "checkout_comment": checkout_comment,
            }},
        )
        if not result.modified_count:
            flash("This visitor has already checked out.", "error")
            return redirect(url_for("visitor_checkout"))
    except PyMongoError as exc:
        app.logger.exception("MongoDB checkout update failed")
        flash(f"Could not complete check-out: {exc}", "error")
        return redirect(url_for("visitor_checkout"))

    return render_template(
        "checkout_success.html",
        visitor_name=visitor.get("visitor_name", ""),
        company=visitor.get("company", ""),
        check_out=now.strftime("%Y-%m-%d %H:%M"),
    )


# ---------------------------------------------------------
# Admin authentication
# ---------------------------------------------------------
@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        password = request.form.get("password", "")

        if password == ADMIN_PASSWORD:
            session["admin"] = True
            return redirect(url_for("admin"))

        flash("Incorrect password.", "error")

    return render_template("login.html", login_error=request.method == "POST")


@app.route("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))


# ---------------------------------------------------------
# Admin visitor list
# ---------------------------------------------------------
@app.route("/admin")
def admin():
    if not admin_required():
        return redirect(url_for("admin_login"))

    q = request.args.get("q", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    status = request.args.get("status", "").strip()

    mongo_filter = build_admin_filter(q, date_from, date_to, status)

    try:
        docs = list(
            visitors_collection.find(mongo_filter)
            .sort("check_in", DESCENDING)
            .limit(1000)
        )
    except PyMongoError as exc:
        app.logger.exception("MongoDB admin query failed")
        flash(f"Could not load visitor records: {exc}", "error")
        docs = []

    visitors = [serialize_visitor(doc) for doc in docs]

    return render_template(
        "admin.html",
        visitors=visitors,
        q=q,
        date_from=date_from,
        date_to=date_to,
        status=status,
    )


@app.post("/admin/checkout/<visitor_id>")
def admin_checkout(visitor_id):
    if not admin_required():
        return redirect(url_for("admin_login"))

    object_id = to_object_id(visitor_id)

    if not object_id:
        flash("Invalid visitor record.", "error")
        return redirect(url_for("admin"))

    try:
        result = visitors_collection.update_one(
            {"_id": object_id, "check_out": None},
            {"$set": {"check_out": datetime.now()}},
        )

        if result.modified_count:
            flash("Visitor checked out successfully.", "success")
        else:
            flash("Visitor is already checked out or was not found.", "error")

    except PyMongoError as exc:
        app.logger.exception("MongoDB admin checkout failed")
        flash(f"Could not complete check-out: {exc}", "error")

    return redirect(request.referrer or url_for("admin"))


# ---------------------------------------------------------
# Excel export
# ---------------------------------------------------------
@app.route("/admin/export")
def export_excel():
    if not admin_required():
        return redirect(url_for("admin_login"))

    q = request.args.get("q", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    status = request.args.get("status", "").strip()

    mongo_filter = build_admin_filter(q, date_from, date_to, status)

    try:
        docs = list(visitors_collection.find(mongo_filter).sort("check_in", DESCENDING))
    except PyMongoError as exc:
        app.logger.exception("MongoDB export query failed")
        flash(f"Could not export visitor records: {exc}", "error")
        return redirect(url_for("admin"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Visitors"

    headers = [
        "ID",
        "Company",
        "Visitor Name",
        "Host Name",
        "Host Phone",
        "Purpose",
        "Comment",
        "Check Out Comment / Issue",
        "Safety Agreed",
        "Check In Date",
        "Check In Time",
        "Check Out Date",
        "Check Out Time",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003A70")
        cell.alignment = Alignment(horizontal="center")

    for doc in docs:
        ws.append(
            [
                str(doc["_id"]),
                doc.get("company", ""),
                doc.get("visitor_name", ""),
                doc.get("host_name", ""),
                doc.get("host_phone", ""),
                doc.get("purpose", ""),
                doc.get("comment", ""),
                doc.get("checkout_comment", ""),
                "Yes" if doc.get("safety_agreed") else "No",
                export_date(doc.get("check_in")),
                export_time(doc.get("check_in")),
                export_date(doc.get("check_out")),
                export_time(doc.get("check_out")),
            ]
        )

    widths = {
        "A": 26,
        "B": 24,
        "C": 22,
        "D": 22,
        "E": 18,
        "F": 24,
        "G": 36,
        "H": 40,
        "I": 16,
        "J": 16,
        "K": 16,
        "L": 16,
        "M": 16,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"visitor_list_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------
# Start application
# ---------------------------------------------------------
if __name__ == "__main__":
    try:
        verify_mongodb()
        ensure_indexes()
        print("MongoDB connection: OK")
        print(f"Database: {MONGO_DB_NAME}")
        print(f"Collection: {MONGO_COLLECTION_NAME}")
    except Exception as exc:
        print("\nMongoDB connection failed.")
        print(exc)
        print("\nCheck MONGO_URI in your .env file and Atlas Network Access settings.")
        raise

    app.run(host="0.0.0.0", port=5000, debug=True)
